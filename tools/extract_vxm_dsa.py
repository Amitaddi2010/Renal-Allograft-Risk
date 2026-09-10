"""
extract_vxm_dsa.py — pull demographics and DSA/MFI out of the VXM workbook.

    python tools/extract_vxm_dsa.py                         # default paths
    python tools/extract_vxm_dsa.py --src <xlsx> --out <csv>

Source: Dataset/Data VXM current.xlsx, one sheet per year.

Layout (verified against the 2025 and 2026 sheets)
--------------------------------------------------
Row 2 carries locus group headers, row 3 the real column names, data from row 4.

    0 S.No   1 Lab No   2 Patient Name   3 Age   4 Sex   5 Blood Gp   6 Pt HLA
    7/8 CR No   9/10 Donor Name   10/11 Age   11/12 Sex   12/13 Blood Gp
    14 CR No   15 Relation   16 HLA Done   17 Lab Id   18 Method

then, for each locus, two (allele, MFI) pairs:

    A*    19,20   21,22        B*    23,24   25,26
    C*    27,28   29,30        DRB1* 31,32   33,34
    DPB1* 35,36   37,38        DQA1* 39,40   41,42
    DQB1* 43,44   45,46

    47 Patient Report   48 Comments   49/50 Class I / Class II epitopes

Column positions drift by a place or two between years (the donor block starts at
9 in 2025 and 10 in 2026), so the header row is read per sheet and columns are
located by name rather than assumed.

Output: one row per patient with demographics plus a long-form DSA list, and a
per-locus wide form. 'NA' and blank MFI cells are preserved as empty, never as 0
- a missing MFI is not a negative result.
"""
import argparse
import csv
import os
import re
import sys
import warnings

warnings.filterwarnings('ignore')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_SRC = os.path.join(os.path.dirname(ROOT), 'Dataset', 'Data VXM current.xlsx')
DEFAULT_OUT = os.path.join(os.path.dirname(ROOT), 'processed_data', 'eplet_batch', 'vxm_dsa.csv')

LOCI = ['A', 'B', 'C', 'DRB1', 'DPB1', 'DQA1', 'DQB1']


def clean(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s.upper() in ('NA', 'N/A', 'NIL', '-', '--') else s


def as_number(v):
    """MFI as a float, or '' when absent. Never coerces a missing value to zero."""
    s = clean(v)
    if not s:
        return ''
    s = s.replace(',', '')
    m = re.match(r'^[<>~]?\s*(\d+(?:\.\d+)?)', s)
    return float(m.group(1)) if m else ''


def header_map(row):
    """column index -> normalised header name, from the row-3 sub-headers."""
    out = {}
    for i, v in enumerate(row):
        s = clean(v)
        if s:
            out[i] = s
    return out


def locate(headers):
    """Find the (allele, MFI) column pairs for each locus, by header name."""
    pairs = {}
    idxs = sorted(headers)
    for i in idxs:
        name = headers[i].rstrip('*').upper()
        if name not in [l.upper() for l in LOCI]:
            continue
        locus = next(l for l in LOCI if l.upper() == name)
        nxt = headers.get(i + 1, '')
        if nxt.upper().startswith('MFI'):
            pairs.setdefault(locus, []).append((i, i + 1))
    return pairs


def field(headers, *names):
    """First column index whose header matches any of the given names."""
    for i in sorted(headers):
        h = headers[i].lower().rstrip('*').strip()
        for n in names:
            if h == n.lower():
                return i
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=DEFAULT_SRC)
    ap.add_argument('--out', default=DEFAULT_OUT)
    ap.add_argument('--mfi-threshold', type=float, default=1000.0,
                    help='MFI at or above which a specificity is counted as positive')
    args = ap.parse_args()

    import openpyxl
    if not os.path.exists(args.src):
        sys.exit('missing %s' % args.src)
    wb = openpyxl.load_workbook(args.src, read_only=True, data_only=True)

    rows, long_rows = [], []
    for sheet in wb.sheetnames:
        if not re.match(r'^20\d\d$', sheet):
            continue
        ws = wb[sheet]
        data = list(ws.iter_rows(values_only=True))
        if len(data) < 4:
            continue
        # The header row moved between years - row 3 in 2024-2026, row 4 in the
        # earlier sheets - so find it rather than assume it. The row carrying
        # 'Patient Name' is the header; data starts on the row after.
        hrow = None
        for i in range(min(8, len(data))):
            cells = [clean(c).lower() for c in (data[i] or ())]
            if 'patient name' in cells:
                hrow = i
                break
        if hrow is None:
            print('  %-6s no header row found, skipped' % sheet)
            continue
        headers = header_map(data[hrow])
        pairs = locate(headers)
        if not pairs:
            print('  %-6s no locus/MFI columns found, skipped' % sheet)
            continue

        c_name = field(headers, 'Patient Name')
        c_age = field(headers, 'Age')
        c_sex = field(headers, 'Sex')
        c_bg = field(headers, 'Blood Gp')
        c_rel = field(headers, 'Relation')
        c_lab = field(headers, 'Lab No', 'Lab Id')
        # the donor block repeats Age/Sex/Blood Gp after the donor name
        c_dname = field(headers, 'Donor Name')
        def after(idx, *names):
            if idx is None:
                return None
            for i in sorted(headers):
                if i <= idx:
                    continue
                h = headers[i].lower().rstrip('*').strip()
                if h in [n.lower() for n in names]:
                    return i
            return None
        c_dage = after(c_dname, 'Age')
        c_dsex = after(c_dname, 'Sex')
        c_dbg = after(c_dname, 'Blood Gp')

        n_sheet = 0
        for r in data[hrow + 1:]:
            if not r:
                continue
            name = clean(r[c_name]) if c_name is not None and c_name < len(r) else ''
            if not name or name.lower() in ('patient name', 's.no'):
                continue
            rec = {
                'year': sheet,
                'lab_no': clean(r[c_lab]) if c_lab is not None and c_lab < len(r) else '',
                'patient': name,
                'pt_age': clean(r[c_age]) if c_age is not None and c_age < len(r) else '',
                'pt_sex': clean(r[c_sex]) if c_sex is not None and c_sex < len(r) else '',
                'pt_blood_group': clean(r[c_bg]) if c_bg is not None and c_bg < len(r) else '',
                'donor': clean(r[c_dname]) if c_dname is not None and c_dname < len(r) else '',
                'dn_age': clean(r[c_dage]) if c_dage is not None and c_dage < len(r) else '',
                'dn_sex': clean(r[c_dsex]) if c_dsex is not None and c_dsex < len(r) else '',
                'dn_blood_group': clean(r[c_dbg]) if c_dbg is not None and c_dbg < len(r) else '',
                'relation': clean(r[c_rel]) if c_rel is not None and c_rel < len(r) else '',
            }

            mfis, specs, above = [], 0, 0
            for locus, cols in pairs.items():
                for k, (ca, cm) in enumerate(cols, start=1):
                    allele = clean(r[ca]) if ca < len(r) else ''
                    mfi = as_number(r[cm]) if cm < len(r) else ''
                    rec['%s_%d' % (locus, k)] = allele
                    rec['%s_%d_MFI' % (locus, k)] = mfi
                    if allele:
                        specs += 1
                        long_rows.append({'year': sheet, 'lab_no': rec['lab_no'],
                                          'patient': name, 'locus': locus,
                                          'allele': allele, 'mfi': mfi})
                    if mfi != '':
                        mfis.append(mfi)
                        if mfi >= args.mfi_threshold:
                            above += 1
            rec['n_specificities'] = specs
            rec['n_with_mfi'] = len(mfis)
            rec['peak_mfi'] = max(mfis) if mfis else ''
            rec['sum_mfi'] = round(sum(mfis), 1) if mfis else ''
            rec['n_mfi_ge_threshold'] = above if mfis else ''
            rows.append(rec)
            n_sheet += 1
        print('  %-6s %5d patients | %d locus columns' % (sheet, n_sheet, sum(len(v) for v in pairs.values())))

    wb.close()
    if not rows:
        sys.exit('no rows extracted')

    cols = ['year', 'lab_no', 'patient', 'pt_age', 'pt_sex', 'pt_blood_group',
            'donor', 'dn_age', 'dn_sex', 'dn_blood_group', 'relation',
            'n_specificities', 'n_with_mfi', 'peak_mfi', 'sum_mfi', 'n_mfi_ge_threshold']
    for locus in LOCI:
        for k in (1, 2):
            cols += ['%s_%d' % (locus, k), '%s_%d_MFI' % (locus, k)]
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, 'w', newline='', encoding='utf-8') as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)

    long_out = args.out.replace('.csv', '_long.csv')
    with open(long_out, 'w', newline='', encoding='utf-8') as fh:
        w = csv.DictWriter(fh, fieldnames=['year', 'lab_no', 'patient', 'locus', 'allele', 'mfi'])
        w.writeheader()
        w.writerows(long_rows)

    with_mfi = [r for r in rows if r['n_with_mfi']]
    print('\n%d patients -> %s' % (len(rows), args.out))
    print('%d specificity rows -> %s' % (len(long_rows), long_out))
    print('patients with at least one numeric MFI: %d (%.1f%%)'
          % (len(with_mfi), 100.0 * len(with_mfi) / len(rows)))
    if with_mfi:
        peaks = sorted(r['peak_mfi'] for r in with_mfi)
        print('peak MFI: median %.0f, range %.0f-%.0f'
              % (peaks[len(peaks) // 2], peaks[0], peaks[-1]))
    return 0


if __name__ == '__main__':
    sys.exit(main())
